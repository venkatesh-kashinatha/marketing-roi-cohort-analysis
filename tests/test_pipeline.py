"""Tests run the real SQL on a tiny dataset whose answers can be worked out by hand."""

import textwrap

import pandas as pd
import pytest
from openpyxl import load_workbook

from mroi.database import build_database, params, run_checks
from mroi.generate import generate
from mroi.insights import CUT, OPTIMIZE, SCALE, TOO_EARLY, action, scenario

FIXTURE = {
    "campaigns.csv": """
        campaign_id,campaign_name,channel,objective,start_date,end_date
        A,Search A,Paid Search,Prospecting,2024-01-01,2024-12-31
        B,Display B,Display,Awareness,2024-01-01,2024-12-31
    """,
    "daily_spend.csv": """
        spend_date,campaign_id,spend,impressions,clicks
        2024-01-01,A,150,1500,75
        2024-01-01,B,200,10000,400
        2024-06-30,A,50,500,25
    """,
    "signups.csv": """
        user_id,signup_date,campaign_id
        U1,2024-01-01,A
        U2,2024-01-01,A
        U3,2024-01-01,A
        U4,2024-01-01,B
        U5,2024-06-20,A
    """,
    # days after sign-up: O1 2, O2 45, O3 35, O4 91, O5 9, O6 5
    "orders.csv": """
        order_id,user_id,order_date,revenue,cost_of_goods
        O1,U1,2024-01-03,100,40
        O2,U1,2024-02-15,50,20
        O3,U2,2024-02-05,80,30
        O4,U1,2024-04-01,70,30
        O5,U4,2024-01-10,40,25
        O6,U5,2024-06-25,60,20
    """,
}


def write_fixture(folder, extra_orders=""):
    for name, text in FIXTURE.items():
        body = textwrap.dedent(text).strip() + "\n"
        if name == "orders.csv" and extra_orders:
            body += extra_orders + "\n"
        (folder / name).write_text(body)
    return folder


@pytest.fixture()
def con(tmp_path):
    connection = build_database(write_fixture(tmp_path), ":memory:", horizon=60)
    yield connection
    connection.close()


def test_params_and_mature_cutoff(con):
    p = params(con)
    assert p["end_date"] == pd.Timestamp("2024-06-30")
    assert p["mature_cutoff"] == pd.Timestamp("2024-05-02")  # 60 days of history: Jun 30 minus 59 days


def test_users_flags_and_horizon_value(con):
    u = con.execute("SELECT * FROM int_users ORDER BY user_id").df().set_index("user_id")
    assert list(u["converted_7d"]) == [1, 0, 0, 0, 1]
    assert list(u["converted_30d"]) == [1, 0, 0, 1, 1]
    assert list(u["converted_in_horizon"]) == [1, 1, 0, 1, 1]
    assert u.loc["U1", "orders_in_horizon"] == 2            # the day-91 order is outside 60 days
    assert u.loc["U1", "gross_profit_in_horizon"] == pytest.approx(90)
    assert list(u["is_mature"]) == [1, 1, 1, 1, 0]          # U5 signed up after the cutoff


def test_scorecard_matches_hand_calculation(con):
    sc = con.execute("SELECT * FROM mart_campaign_scorecard").df().set_index("campaign_id")
    a, b = sc.loc["A"], sc.loc["B"]
    assert a["spend"] == pytest.approx(150)                 # the Jun 30 spend is after the cutoff
    assert (a["signups"], a["customers"], a["customers_7d"]) == (3, 2, 1)
    assert a["revenue"] == pytest.approx(230)
    assert a["gross_profit"] == pytest.approx(140)
    assert a["cac"] == pytest.approx(75)
    assert a["roas"] == pytest.approx(230 / 150)
    assert a["roi"] == pytest.approx((140 - 150) / 150)
    assert a["payback_months"] == 4                         # 180 of gross profit covers 150 by day 119
    assert a["current_annual_spend"] == pytest.approx(200)
    assert b["cac"] == pytest.approx(200)
    assert b["roi"] == pytest.approx((15 - 200) / 200)
    assert pd.isna(b["payback_months"])


def test_payback_counts_only_full_windows(con):
    pb = con.execute("SELECT * FROM mart_payback WHERE campaign_id = 'A' ORDER BY period").df()
    assert list(pb["period"]) == [0, 1, 2, 3, 4, 5]          # period 6 needs sign-ups before 2024
    assert list(pb["gross_profit_cum"]) == pytest.approx([60, 140, 140, 180, 180, 180])
    assert pb.loc[0, "signups"] == 3                         # U5 lacks 30 days of history
    assert pb.loc[3, "payback_ratio"] == pytest.approx(1.2)


def test_cohort_conversion_maturity(con):
    cc = con.execute("SELECT * FROM mart_cohort_conversion ORDER BY channel, period").df()
    search = cc[cc["channel"] == "Paid Search"]
    assert list(search["period"]) == [0, 1, 2, 3, 4]         # Jan 31 + 179 days is past the data
    assert list(search["customers_cum"]) == [1, 2, 2, 2, 2]
    assert set(cc["signups"]) == {3, 1}
    assert pd.Timestamp("2024-06-01") not in set(pd.to_datetime(cc["cohort_month"]))


def test_channel_monthly_keeps_months_with_only_first_purchases(con):
    m = con.execute("SELECT * FROM mart_channel_monthly WHERE channel = 'Paid Search' ORDER BY month").df()
    feb = m[pd.to_datetime(m["month"]) == "2024-02-01"]
    assert feb["new_customers"].item() == 1 and feb["spend"].item() == 0


def test_checks_pass_and_catch_problems(con, tmp_path):
    assert run_checks(con)["failing_rows"].sum() == 0
    broken_dir = tmp_path / "broken"
    broken_dir.mkdir()
    broken = build_database(write_fixture(broken_dir, "O7,U3,2023-12-31,10,5"), ":memory:", horizon=60)
    checks = run_checks(broken).set_index("check_name")["failing_rows"]
    assert checks["Orders dated before the sign-up"] == 1


def test_action_rules():
    base = {"signups": 1000, "spend": 1000.0, "payback_months": float("nan")}
    assert action(pd.Series({**base, "roi": 0.8})) == SCALE
    assert action(pd.Series({**base, "roi": 0.2})) == OPTIMIZE
    assert action(pd.Series({**base, "roi": -0.1, "payback_months": 11})) == OPTIMIZE
    assert action(pd.Series({**base, "roi": -0.1})) == CUT
    assert action(pd.Series({**base, "signups": 50, "roi": 2.0})) == TOO_EARLY


def test_scenario_math():
    sc = pd.DataFrame({"current_annual_spend": [1000.0, 2000.0], "cac": [50.0, 100.0],
                       "gp_per_customer": [100.0, 50.0]})
    s = scenario(sc, pd.Series([0.5, -0.5]))
    assert s["spend_change"] == pytest.approx(500 - 1000)
    assert s["extra_customers"] == pytest.approx(10 - 10)
    assert s["net_gain"] == pytest.approx((1000 - 500) + (-500 + 1000))


def test_generator_is_reproducible():
    a = generate("2024-01-01", "2024-03-31", seed=7)
    b = generate("2024-01-01", "2024-03-31", seed=7)
    assert all(a[k].equals(b[k]) for k in a)
    orders = a["orders"].merge(a["signups"], on="user_id")
    assert (pd.to_datetime(orders["order_date"]) >= pd.to_datetime(orders["signup_date"])).all()


def test_full_build_writes_report(tmp_path):
    from mroi.__main__ import main

    raw = tmp_path / "raw"
    raw.mkdir()
    write_fixture(raw)
    assert main(["build", "--data", str(raw), "--out", str(tmp_path / "out"), "--horizon", "60"]) == 0
    wb = load_workbook(tmp_path / "out" / "Marketing_ROI_Report.xlsx")
    assert wb.sheetnames == ["Summary", "Scorecard", "Cohort Conversion", "Payback", "Budget Scenario",
                             "Channel Monthly", "Data Checks", "Definitions"]
    assert (tmp_path / "out" / "marts" / "fact_users.csv").exists()
