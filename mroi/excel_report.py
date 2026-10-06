"""Write the Excel report.

SQL supplies the counts and sums. Every rate (conversion, CAC, ROAS, ROI, payback) is an Excel
formula on top of them, and the campaign actions follow decision rules held in input cells on
the Summary sheet, so the workbook can be audited and the rules changed without rerunning SQL.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from . import __version__
from .config import CUT_ROI, MIN_SIGNUPS, SCALE_ROI
from .insights import day

FONT = "Arial"
NAVY = "1F3A5F"
TEAL = "1D6F6A"
BAND = "EEF2F7"
GREY = "595959"
INPUT_FILL = "FFFF00"
INPUT_FONT = "0000FF"
GOOD, WARN, BAD = "C6EFCE", "FFEB9C", "FFC7CE"

MONEY = '$#,##0;($#,##0);"-"'
MONEY2 = '$#,##0.00;($#,##0.00);"-"'
INT = '#,##0;(#,##0);"-"'
PCT0 = '0%;(0%);"0%"'
PCT1 = '0.0%;(0.0%);"-"'
PCT2 = '0.00%;(0.00%);"-"'
TIMES = '0.00"x";(0.00"x");"-"'
MONTH = "mmm yyyy"
THIN = Side(style="thin", color="8EA0B8")
PERIODS = range(12)


@dataclass
class ReportData:
    scorecard: pd.DataFrame          # mart_campaign_scorecard + action, sorted for display
    cohort_signups: pd.DataFrame     # cohort_month, channel, signups (all cohorts)
    cohorts: pd.DataFrame            # mart_cohort_conversion
    payback: pd.DataFrame            # mart_payback
    monthly: pd.DataFrame            # mart_channel_monthly
    checks: pd.DataFrame             # check_name, failing_rows
    changes: pd.Series               # default budget changes, aligned with scorecard rows
    findings: list[str]
    params: dict
    counts: dict
    source: str


def write_report(path: Path, d: ReportData) -> Path:
    wb = Workbook()
    _arial_default(wb)
    summary = wb.active
    summary.title = "Summary"
    sheets = {name: wb.create_sheet(name) for name in
              ("Scorecard", "Cohort Conversion", "Payback", "Budget Scenario", "Channel Monthly",
               "Data Checks", "Definitions")}

    _summary_inputs(wb, summary, d)
    rows = _scorecard(sheets["Scorecard"], d)
    _summary_outputs(summary, d, rows)
    _cohorts(sheets["Cohort Conversion"], d)
    _payback(sheets["Payback"], d)
    _budget(sheets["Budget Scenario"], d, rows)
    _monthly(sheets["Channel Monthly"], d)
    _checks(sheets["Data Checks"], d)
    _definitions(sheets["Definitions"], d)

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "portrait" if ws is summary else "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return path


# --------------------------------------------------------------------------- helpers

def _arial_default(wb: Workbook) -> None:
    try:
        default = Font(name=FONT, size=10)
        wb._fonts[0] = default
        wb._named_styles["Normal"].font = default
    except (AttributeError, IndexError, KeyError, TypeError):
        pass


def _font(bold=False, color="000000", size=10, italic=False) -> Font:
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def _define(wb: Workbook, name: str, ref: str) -> None:
    dn = DefinedName(name, attr_text=ref)
    try:
        wb.defined_names[name] = dn
    except TypeError:
        wb.defined_names.append(dn)


def _title(ws, text: str, subtitle: str | None = None) -> None:
    ws["A1"] = text
    ws["A1"].font = _font(bold=True, color=NAVY, size=14)
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = _font(italic=True, color=GREY, size=9)


def _header(ws, row: int, labels: list[str], col: int = 1, fill: str = NAVY) -> None:
    for i, label in enumerate(labels):
        cell = ws.cell(row=row, column=col + i, value=label)
        cell.font = _font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def _section(ws, row: int, text: str, width: int, col: int = 1) -> None:
    for c in range(col, col + width):
        cell = ws.cell(row=row, column=c)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = _font(bold=True, color="FFFFFF")
    ws.cell(row=row, column=col, value=text)


def _label(ws, row: int, text: str) -> None:
    ws.cell(row=row, column=1, value=text).font = _font(bold=True, color=NAVY, size=11)


def _put(ws, row: int, col: int, value, fmt: str | None = None, bold: bool = False):
    cell = ws.cell(row=row, column=col, value=_clean(value))
    if fmt:
        cell.number_format = fmt
    if bold:
        cell.font = _font(bold=True)
    return cell


def _clean(v):
    if v is None or (not isinstance(v, str) and pd.isna(v)):
        return None
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    if hasattr(v, "item"):
        return v.item()
    return v


def _input(cell, fmt: str) -> None:
    cell.number_format = fmt
    cell.font = _font(bold=True, color=INPUT_FONT)
    cell.fill = PatternFill("solid", fgColor=INPUT_FILL)
    cell.alignment = Alignment(horizontal="right")


def _total_row(ws, row: int, first_col: int, last_col: int) -> None:
    for c in range(first_col, last_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = _font(bold=True)
        cell.fill = PatternFill("solid", fgColor=BAND)
        cell.border = Border(top=THIN)


def _roi_colors(ws, cell_range: str) -> None:
    """Green at or above the scale threshold, red below the cut threshold, amber between."""
    first = cell_range.split(":")[0].replace("$", "")
    ws.conditional_formatting.add(cell_range, FormulaRule(
        formula=[f"AND(ISNUMBER({first}),{first}>=ScaleROI)"], fill=PatternFill("solid", fgColor=GOOD)))
    ws.conditional_formatting.add(cell_range, FormulaRule(
        formula=[f"AND(ISNUMBER({first}),{first}<CutROI)"], fill=PatternFill("solid", fgColor=BAD)))
    ws.conditional_formatting.add(cell_range, FormulaRule(
        formula=[f"AND(ISNUMBER({first}),{first}>=CutROI,{first}<ScaleROI)"],
        fill=PatternFill("solid", fgColor=WARN)))


def _action_colors(ws, cell_range: str) -> None:
    first = cell_range.split(":")[0].replace("$", "")
    for text, color in (("Scale", "006100"), ("Cut or rework", "9C0006"), ("Optimize", "7F6000")):
        ws.conditional_formatting.add(cell_range, FormulaRule(
            formula=[f'{first}="{text}"'], font=Font(name=FONT, bold=True, color=color)))


def _widths(ws, widths: dict[str, float]) -> None:
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width


def _wrap_lines(ws, start_row: int, lines: list[str], width_cols: int, chars_per_line: int) -> int:
    for k, text in enumerate(lines):
        r = start_row + k
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width_cols)
        cell = ws.cell(row=r, column=1, value=f"• {text}")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 13.5 * (1 + (len(text) + 2) // chars_per_line)
    return start_row + len(lines)


def _period_labels() -> list[str]:
    return [f"{30 * (k + 1)} days" for k in PERIODS]


# --------------------------------------------------------------------------- Summary

def _summary_inputs(wb: Workbook, ws, d: ReportData) -> None:
    p = d.params
    _title(ws, "Marketing ROI and Cohort Analysis",
           f"Sign-ups {day(p['first_signup'])} to {day(p['mature_cutoff'])}, outcomes within "
           f"{p['horizon']} days · built {datetime.now():%b %d, %Y %H:%M} by mroi v{__version__}")
    _section(ws, 4, "Decision rules (change the yellow cells)", 2)
    rules = [
        (f"Scale if {p['horizon']}-day ROI is at least", SCALE_ROI, PCT0, "ScaleROI"),
        ("Cut if ROI is below, with no 12-month payback", CUT_ROI, PCT0, "CutROI"),
        ("Minimum sign-ups to judge a campaign", MIN_SIGNUPS, INT, "MinSignups"),
    ]
    for r, (label, value, fmt, name) in enumerate(rules, start=5):
        ws.cell(row=r, column=1, value=label)
        _input(ws.cell(row=r, column=2, value=value), fmt)
        _define(wb, name, f"Summary!$B${r}")


def _summary_outputs(ws, d: ReportData, rows: dict) -> None:
    t = rows["total"]
    h = d.params["horizon"]
    _section(ws, 9, "Headline KPIs (measured window)", 2)
    kpis = [
        ("Spend", f"=Scorecard!D{t}", MONEY),
        ("Sign-ups", f"=Scorecard!G{t}", INT),
        (f"Customers within {h} days", f"=Scorecard!I{t}", INT),
        ("Sign-up to customer conversion", f"=Scorecard!T{t}", PCT1),
        ("Blended CAC", f"=Scorecard!U{t}", MONEY2),
        (f"ROAS ({h}-day revenue ÷ spend)", f"=Scorecard!W{t}", TIMES),
        (f"ROI ({h}-day gross profit vs spend)", f"=Scorecard!X{t}", PCT0),
    ]
    for r, (label, formula, fmt) in enumerate(kpis, start=10):
        ws.cell(row=r, column=1, value=label)
        _put(ws, r, 2, formula, fmt, bold=True).alignment = Alignment(horizontal="right")
    _roi_colors(ws, "B16:B16")

    top = 19
    _section(ws, top, "Campaigns", 7)
    _header(ws, top + 1, ["Campaign", "Channel", "Spend", "CAC", "ROI", "Pays back within (months)", "Action"])
    for k, r_sc in enumerate(rows["campaigns"]):
        r = top + 2 + k
        for c, (col, fmt) in enumerate(zip("BCDUXMZ", (None, None, MONEY, MONEY2, PCT0, INT, None)), start=1):
            cell = _put(ws, r, c, f"=Scorecard!{col}{r_sc}", fmt)
            if c >= 3:
                cell.alignment = (Alignment(horizontal="right") if c < 7 else Alignment(horizontal="left", indent=1))
    last = top + 1 + len(rows["campaigns"])
    _roi_colors(ws, f"E{top + 2}:E{last}")
    _action_colors(ws, f"G{top + 2}:G{last}")

    ftop = last + 2
    _section(ws, ftop, "Key findings and recommendation", 7)
    end = _wrap_lines(ws, ftop + 1, d.findings, 7, 135)
    ws.cell(row=end + 1, column=1, value="Findings are written by Python for this run; the tables above update "
                                          "when you change the decision rules.").font = _font(italic=True, color=GREY, size=9)
    _widths(ws, {"A": 44, "B": 15, "C": 12, "D": 11, "E": 9, "F": 13, "G": 19})


# --------------------------------------------------------------------------- Scorecard

def _scorecard(ws, d: ReportData) -> dict:
    h = d.params["horizon"]
    p = d.params
    _title(ws, "Campaign scorecard",
           f"Sign-ups (and the spend that bought them) from {day(p['first_signup'])} to "
           f"{day(p['mature_cutoff'])}, with every outcome counted within {h} days of sign-up.")
    sql_cols = [
        ("Campaign ID", "campaign_id", None), ("Campaign", "campaign_name", None), ("Channel", "channel", None),
        ("Spend", "spend", MONEY), ("Impressions", "impressions", INT), ("Clicks", "clicks", INT),
        ("Sign-ups", "signups", INT), ("Customers in 7 days", "customers_7d", INT),
        (f"Customers in {h} days", "customers", INT), (f"Orders in {h} days", "orders", INT),
        (f"Revenue in {h} days", "revenue", MONEY), (f"Gross profit in {h} days", "gross_profit", MONEY),
        ("Pays back within (months)", "payback_months", INT),
        ("Spend, last 12 months (live campaigns)", "current_annual_spend", MONEY),
    ]
    calc_cols = [
        ("CTR", '=IF(E{r}=0,"",F{r}/E{r})', PCT2),
        ("CPC", '=IF(F{r}=0,"",D{r}/F{r})', MONEY2),
        ("Sign-up rate", '=IF(F{r}=0,"",G{r}/F{r})', PCT1),
        ("Cost per sign-up", '=IF(G{r}=0,"",D{r}/G{r})', MONEY2),
        ("7-day conversion", '=IF(G{r}=0,"",H{r}/G{r})', PCT1),
        (f"{h}-day conversion", '=IF(G{r}=0,"",I{r}/G{r})', PCT1),
        ("CAC", '=IF(I{r}=0,"",D{r}/I{r})', MONEY2),
        ("Gross profit per customer", '=IF(I{r}=0,"",L{r}/I{r})', MONEY2),
        ("ROAS", '=IF(D{r}=0,"",K{r}/D{r})', TIMES),
        ("ROI", '=IF(D{r}=0,"",(L{r}-D{r})/D{r})', PCT0),
        ("LTV:CAC", '=IF(OR(U{r}="",V{r}=""),"",V{r}/U{r})', TIMES),
        ("Action", '=IF(OR(G{r}<MinSignups,X{r}=""),"Too early to judge",IF(X{r}>=ScaleROI,"Scale",'
                   'IF(OR(X{r}>=CutROI,ISNUMBER(M{r})),"Optimize","Cut or rework")))', None),
    ]
    n_sql = len(sql_cols)
    last_col = n_sql + len(calc_cols)
    ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=n_sql)
    ws.merge_cells(start_row=4, start_column=n_sql + 1, end_row=4, end_column=last_col)
    for col, text, fill in ((1, "From SQL", NAVY), (n_sql + 1, "Calculated in Excel", TEAL)):
        cell = ws.cell(row=4, column=col, value=text)
        cell.font = _font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(horizontal="center")
    _header(ws, 5, [c[0] for c in sql_cols])
    _header(ws, 5, [c[0] for c in calc_cols], col=n_sql + 1, fill=TEAL)
    ws.row_dimensions[5].height = 42

    first = 6
    campaign_rows = []
    for i, row in enumerate(d.scorecard.itertuples(index=False)):
        r = first + i
        campaign_rows.append(r)
        for c, (_, field, fmt) in enumerate(sql_cols, start=1):
            value = getattr(row, field)
            if field == "payback_months" and pd.isna(value):
                value = "Not yet"
            _put(ws, r, c, value, fmt)
        for c, (_, formula, fmt) in enumerate(calc_cols, start=n_sql + 1):
            _put(ws, r, c, formula.format(r=r), fmt)
    total = first + len(d.scorecard)
    ws.cell(row=total, column=2, value="Total")
    for c in list(range(4, 13)) + [14]:
        letter = get_column_letter(c)
        _put(ws, total, c, f"=SUM({letter}{first}:{letter}{total - 1})", sql_cols[c - 1][2])
    for c, (_, formula, fmt) in enumerate(calc_cols[:-1], start=n_sql + 1):
        _put(ws, total, c, formula.format(r=total), fmt)
    _total_row(ws, total, 1, last_col)
    for r in range(first, total + 1):
        ws.cell(row=r, column=13).alignment = Alignment(horizontal="right")

    _roi_colors(ws, f"X{first}:X{total}")
    _action_colors(ws, f"Z{first}:Z{total - 1}")
    ws.conditional_formatting.add(f"D{first}:D{total - 1}", DataBarRule(
        start_type="num", start_value=0, end_type="max", color="9DB9E0"))
    ws.freeze_panes = f"C{first}"
    _widths(ws, {"A": 11, "B": 30, "C": 13, **{get_column_letter(c): 12 for c in range(4, last_col + 1)},
                 "N": 14, "Z": 18})
    return {"campaigns": campaign_rows, "total": total}


# --------------------------------------------------------------------------- Cohorts

def _cohorts(ws, d: ReportData) -> None:
    _title(ws, "Cohort conversion",
           "Each row is the month people signed up. Columns show how many had made a first purchase within "
           "30, 60, ... 360 days. A cell stays blank until every member of the cohort has had that long.")
    signups = d.cohort_signups.groupby("cohort_month")["signups"].sum().sort_index()
    customers = (d.cohorts.groupby(["cohort_month", "period"])["customers_cum"].sum()
                 .unstack("period").reindex(signups.index))
    n = len(signups)
    labels = _period_labels()

    _label(ws, 4, "Customers who had bought (cumulative count)")
    _header(ws, 5, ["Sign-up month", "Sign-ups", *labels])
    first1 = 6
    for i, (month, count) in enumerate(signups.items()):
        r = first1 + i
        _put(ws, r, 1, month, MONTH)
        _put(ws, r, 2, count, INT)
        for k in PERIODS:
            value = customers.loc[month, k] if k in customers.columns else None
            _put(ws, r, 3 + k, value, INT)
    last1 = first1 + n - 1

    top2 = last1 + 3
    _label(ws, top2, "Conversion rate (cumulative % of sign-ups who had bought)")
    _header(ws, top2 + 1, ["Sign-up month", "Sign-ups", *labels])
    first2 = top2 + 2
    for i in range(n):
        r, src = first2 + i, first1 + i
        _put(ws, r, 1, f"=A{src}", MONTH)
        _put(ws, r, 2, f"=B{src}", INT)
        for k in PERIODS:
            col = get_column_letter(3 + k)
            _put(ws, r, 3 + k, f'=IF({col}{src}="","",{col}{src}/$B{src})', PCT1)
    last2 = first2 + n - 1
    ws.conditional_formatting.add(f"C{first2}:N{last2}", ColorScaleRule(
        start_type="min", start_color="FFFFFF", end_type="max", end_color="5B8DD6"))

    chart = LineChart()
    chart.title = "30-day conversion by sign-up month"
    chart.height, chart.width = 8, 18
    chart.y_axis.number_format = "0%"
    chart.x_axis.number_format = "mmm yy"
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.legend = None
    chart.add_data(Reference(ws, min_col=3, min_row=top2 + 1, max_row=last2), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=first2, max_row=last2))
    for s in chart.series:
        s.smooth = False
    ws.add_chart(chart, "P4")
    ws.freeze_panes = "C6"
    _widths(ws, {"A": 14, "B": 10, **{get_column_letter(c): 9 for c in range(3, 15)}})


# --------------------------------------------------------------------------- Payback

def _payback(ws, d: ReportData) -> None:
    _title(ws, "Payback by channel",
           "Cumulative gross profit from a channel's sign-ups ÷ the spend that bought them. 100% means the spend "
           "has been earned back. Each column only counts sign-ups with that much history.")
    grouped = d.payback.groupby(["channel", "period"])[["spend", "gross_profit_cum"]].sum()
    channels = sorted(d.payback["channel"].unique())
    labels = _period_labels()
    n = len(channels)
    first_p, last_p = "B", get_column_letter(1 + len(labels))  # periods sit in columns B to M

    def block(top: int, title: str, field: str) -> int:
        _label(ws, top, title)
        _header(ws, top + 1, ["Channel", *labels])
        for i, ch in enumerate(channels):
            r = top + 2 + i
            ws.cell(row=r, column=1, value=ch)
            for k in PERIODS:
                _put(ws, r, 2 + k, grouped[field].get((ch, k)), MONEY)
        return top + 2

    first1 = block(4, "Spend on the sign-ups included", "spend")
    first2 = block(first1 + n + 2, "Cumulative gross profit from those sign-ups", "gross_profit_cum")

    top3 = first2 + n + 2
    _label(ws, top3, "Payback ratio (gross profit ÷ spend)")
    _header(ws, top3 + 1, ["Channel", *labels, "Pays back within"])
    first3 = top3 + 2
    top4 = first3 + n + 4
    for i, ch in enumerate(channels):
        r, s_row, g_row, f_row = first3 + i, first1 + i, first2 + i, top4 + 2 + i
        ws.cell(row=r, column=1, value=ch)
        for k in PERIODS:
            col = get_column_letter(2 + k)
            _put(ws, r, 2 + k, f'=IF(OR({col}{s_row}="",{col}{s_row}=0),"",{col}{g_row}/{col}{s_row})', PCT0)
        _put(ws, r, 14, f'=IFERROR(INDEX(${first_p}${top3 + 1}:${last_p}${top3 + 1},'
                        f'MATCH(1,{first_p}{f_row}:{last_p}{f_row},0)),"Not within 12 months")'
             ).alignment = Alignment(horizontal="center")
    breakeven = first3 + n
    ws.cell(row=breakeven, column=1, value="Break-even").font = _font(italic=True, color=GREY)
    for k in PERIODS:
        _put(ws, breakeven, 2 + k, 1, PCT0).font = _font(italic=True, color=GREY)
    ws.conditional_formatting.add(f"{first_p}{first3}:{last_p}{first3 + n - 1}", CellIsRule(
        operator="greaterThanOrEqual", formula=["1"], fill=PatternFill("solid", fgColor=GOOD)))

    _label(ws, top4, "Helper: 1 once the channel has paid back (feeds 'Pays back within')")
    _header(ws, top4 + 1, ["Channel", *labels])
    for i, ch in enumerate(channels):
        r, ratio_row = top4 + 2 + i, first3 + i
        ws.cell(row=r, column=1, value=ch).font = _font(color=GREY)
        for k in PERIODS:
            col = get_column_letter(2 + k)
            _put(ws, r, 2 + k, f"=IF(AND(ISNUMBER({col}{ratio_row}),{col}{ratio_row}>=1),1,0)").font = _font(color=GREY)

    chart = LineChart()
    chart.title = "Payback ratio by channel (100% = paid back)"
    chart.height, chart.width = 9, 20
    chart.y_axis.number_format = "0%"
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.add_data(Reference(ws, min_col=1, min_row=first3, max_col=13, max_row=breakeven),
                   from_rows=True, titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=2, min_row=top3 + 1, max_col=13, max_row=top3 + 1))
    for s in chart.series:
        s.smooth = False
    ws.add_chart(chart, "P4")
    _widths(ws, {"A": 14, **{get_column_letter(c): 10 for c in range(2, 14)}, "N": 20})


# --------------------------------------------------------------------------- Budget scenario

def _budget(ws, d: ReportData, rows: dict) -> None:
    h = d.params["horizon"]
    _title(ws, "Budget scenario",
           "Change the yellow cells to test budget moves. Projections use each campaign's CAC and gross profit per "
           "customer from the Scorecard, over a year of spend.")
    headers = ["Campaign", "Channel", "Action", "Spend, last 12 months", "Change in spend", "Scenario spend",
               "CAC", "Gross profit per customer", "Customers a year (now)", "Customers a year (scenario)",
               f"{h}-day gross profit (now)", f"{h}-day gross profit (scenario)", "Gross profit after spend: change"]
    _header(ws, 4, headers)
    ws.row_dimensions[4].height = 42
    first = 5
    for i, r_sc in enumerate(rows["campaigns"]):
        r = first + i
        _put(ws, r, 1, f"=Scorecard!B{r_sc}")
        _put(ws, r, 2, f"=Scorecard!C{r_sc}")
        _put(ws, r, 3, f"=Scorecard!Z{r_sc}")
        _put(ws, r, 4, f"=Scorecard!N{r_sc}", MONEY)
        _input(ws.cell(row=r, column=5, value=float(d.changes.iloc[i])), PCT0)
        _put(ws, r, 6, f"=D{r}*(1+E{r})", MONEY)
        _put(ws, r, 7, f"=Scorecard!U{r_sc}", MONEY2)
        _put(ws, r, 8, f"=Scorecard!V{r_sc}", MONEY2)
        _put(ws, r, 9, f'=IF(G{r}="","",D{r}/G{r})', INT)
        _put(ws, r, 10, f'=IF(G{r}="","",F{r}/G{r})', INT)
        _put(ws, r, 11, f'=IF(I{r}="","",I{r}*H{r})', MONEY)
        _put(ws, r, 12, f'=IF(J{r}="","",J{r}*H{r})', MONEY)
        _put(ws, r, 13, f'=IF(OR(K{r}="",L{r}=""),"",(L{r}-F{r})-(K{r}-D{r}))', MONEY)
    total = first + len(rows["campaigns"])
    ws.cell(row=total, column=1, value="Total")
    for c in (4, 6, 9, 10, 11, 12, 13):
        letter = get_column_letter(c)
        _put(ws, total, c, f"=SUM({letter}{first}:{letter}{total - 1})", INT if c in (9, 10) else MONEY)
    _total_row(ws, total, 1, 13)
    _action_colors(ws, f"C{first}:C{total - 1}")

    top = total + 2
    _section(ws, top, "What the scenario does", 4)
    effects = [
        ("Change in annual spend", f"=F{total}-D{total}", MONEY),
        ("Extra customers a year", f"=J{total}-I{total}", INT),
        (f"Extra {h}-day gross profit", f"=L{total}-K{total}", MONEY),
        ("Gross profit after spend: change", f"=M{total}", MONEY),
    ]
    for k, (label, formula, fmt) in enumerate(effects, start=1):
        ws.cell(row=top + k, column=1, value=label)
        _put(ws, top + k, 4, formula, fmt, bold=True)
    note = top + len(effects) + 2
    ws.cell(row=note, column=1, value=(
        "Assumes CAC and gross profit per customer stay the same as budgets change. In practice CAC usually rises "
        "as a campaign scales (smaller audiences, higher bids), so treat increases as an upper bound and test them "
        "in steps. Defaults: halve campaigns marked 'Cut or rework', grow those marked 'Scale' by 50%."
    )).font = _font(italic=True, color=GREY, size=9)
    _widths(ws, {"A": 30, "B": 13, "C": 17, **{get_column_letter(c): 14 for c in range(4, 14)}})
    ws.freeze_panes = "B5"


# --------------------------------------------------------------------------- Channel monthly

def _monthly(ws, d: ReportData) -> None:
    _title(ws, "Monthly trend by channel",
           "Spend and first purchases by calendar month. Monthly CAC here is operational: a month's spend divided by "
           "that month's first purchases, some from people who signed up earlier.")
    m = d.monthly
    months = sorted(m["month"].unique())
    channels = sorted(m["channel"].unique())
    nc = len(channels)
    spend = m.pivot_table(index="month", columns="channel", values="spend", aggfunc="sum").reindex(months)
    new = m.pivot_table(index="month", columns="channel", values="new_customers", aggfunc="sum").reindex(months)
    total_col = nc + 2
    tl = get_column_letter(total_col)
    last_ch = get_column_letter(nc + 1)

    def block(top: int, title: str, pivot: pd.DataFrame, fmt: str) -> int:
        _label(ws, top, title)
        _header(ws, top + 1, ["Month", *channels, "Total"])
        for i, month in enumerate(months):
            r = top + 2 + i
            _put(ws, r, 1, month, MONTH)
            for j, ch in enumerate(channels):
                _put(ws, r, 2 + j, pivot.loc[month, ch] if ch in pivot.columns else None, fmt)
            _put(ws, r, total_col, f"=SUM(B{r}:{last_ch}{r})", fmt, bold=True)
        return top + 2

    first1 = block(4, "Spend", spend, MONEY)
    first2 = block(first1 + len(months) + 2, "First purchases (new customers)", new, INT)
    top3 = first2 + len(months) + 2
    _label(ws, top3, "Monthly CAC (spend ÷ first purchases)")
    _header(ws, top3 + 1, ["Month", *channels, "Blended"])
    first3 = top3 + 2
    for i in range(len(months)):
        r, s_row, n_row = first3 + i, first1 + i, first2 + i
        _put(ws, r, 1, f"=A{s_row}", MONTH)
        for c in range(2, total_col + 1):
            col = get_column_letter(c)
            _put(ws, r, c, f'=IF(OR({col}{n_row}="",{col}{n_row}=0),"",{col}{s_row}/{col}{n_row})', MONEY2)
    last1 = first1 + len(months) - 1
    last3 = first3 + len(months) - 1

    bar = BarChart()
    bar.type = "col"
    bar.grouping = "stacked"
    bar.overlap = 100
    bar.title = "Spend by channel"
    bar.height, bar.width = 8, 20
    bar.y_axis.number_format = "$#,##0"
    bar.x_axis.number_format = "mmm yy"
    bar.x_axis.delete = False
    bar.y_axis.delete = False
    bar.add_data(Reference(ws, min_col=2, min_row=first1 - 1, max_col=nc + 1, max_row=last1), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=1, min_row=first1, max_row=last1))
    ws.add_chart(bar, f"{get_column_letter(total_col + 2)}4")

    line = LineChart()
    line.title = "Blended monthly CAC"
    line.height, line.width = 8, 20
    line.y_axis.number_format = "$#,##0"
    line.x_axis.number_format = "mmm yy"
    line.x_axis.delete = False
    line.y_axis.delete = False
    line.legend = None
    line.add_data(Reference(ws, min_col=total_col, min_row=first3 - 1, max_row=last3), titles_from_data=True)
    line.set_categories(Reference(ws, min_col=1, min_row=first3, max_row=last3))
    for s in line.series:
        s.smooth = False
    ws.add_chart(line, f"{get_column_letter(total_col + 2)}22")
    _widths(ws, {"A": 11, **{get_column_letter(c): 12 for c in range(2, total_col + 1)}, tl: 12})
    ws.freeze_panes = "B6"


# --------------------------------------------------------------------------- Checks and definitions

def _checks(ws, d: ReportData) -> None:
    _title(ws, "Data quality checks", "Run on the staging tables before any analysis. Every check should pass.")
    _header(ws, 4, ["Check", "Failing rows", "Result"])
    for i, row in enumerate(d.checks.itertuples(index=False), start=5):
        ws.cell(row=i, column=1, value=row.check_name)
        _put(ws, i, 2, int(row.failing_rows), INT)
        _put(ws, i, 3, f'=IF(B{i}=0,"Pass","Fail")').alignment = Alignment(horizontal="center")
    last = 4 + len(d.checks)
    ws.conditional_formatting.add(f"C5:C{last}", FormulaRule(
        formula=['C5="Pass"'], fill=PatternFill("solid", fgColor=GOOD)))
    ws.conditional_formatting.add(f"C5:C{last}", FormulaRule(
        formula=['C5="Fail"'], fill=PatternFill("solid", fgColor=BAD)))
    _widths(ws, {"A": 40, "B": 14, "C": 12})


def _definitions(ws, d: ReportData) -> None:
    h = d.params["horizon"]
    _title(ws, "Definitions and run details")
    _header(ws, 3, ["Term", "Definition"])
    terms = [
        ("Sign-up", "A person who registered, credited to the campaign that brought them (first-touch attribution)."),
        ("Customer", f"A sign-up who made a first purchase within {h} days of signing up."),
        ("Measured window", f"Sign-ups on or before the mature cutoff, so each has a full {h} days of history. "
                            "Spend is counted over the same dates."),
        ("Cohort", "Everyone who signed up in the same calendar month."),
        ("Cohort conversion", "Share of a cohort that had bought within 30, 60, ... 360 days. Blank until every member "
                              "has had that long."),
        ("CAC", f"Customer acquisition cost: spend ÷ customers acquired within {h} days."),
        ("ROAS", f"Return on ad spend: {h}-day revenue ÷ spend."),
        ("ROI", f"({h}-day gross profit − spend) ÷ spend. Uses gross profit, not revenue, because revenue ignores "
                "the cost of the goods sold."),
        ("Payback", "The first 30-day period in which cumulative gross profit from a campaign's sign-ups covered the "
                    "spend that bought them."),
        ("LTV:CAC", f"{h}-day gross profit per customer ÷ CAC. Above 1 means a customer is worth more than they cost."),
        ("Actions", "Scale: ROI at or above the scale threshold. Optimize: ROI between the thresholds, or below the cut "
                    "threshold but paying back within 12 months. Cut or rework: below the cut threshold and no "
                    "payback within 12 months. Too early: fewer measured sign-ups than the minimum."),
    ]
    for r, (term, text) in enumerate(terms, start=4):
        ws.cell(row=r, column=1, value=term).font = _font(bold=True)
        ws.cell(row=r, column=2, value=text).alignment = Alignment(wrap_text=True, vertical="top")
    p = d.params
    top = 4 + len(terms) + 1
    _header(ws, top, ["Run detail", "Value"])
    details = [
        ("Source folder", d.source),
        ("Rows loaded", f"{d.counts['stg_campaigns']} campaigns, {d.counts['stg_spend']:,} spend rows, "
                        f"{d.counts['stg_signups']:,} sign-ups, {d.counts['stg_orders']:,} orders"),
        ("Data ends", day(p["end_date"])),
        ("Horizon", f"{h} days"),
        ("Measured window", f"Sign-ups {day(p['first_signup'])} to {day(p['mature_cutoff'])}"),
        ("Built", f"{datetime.now():%Y-%m-%d %H:%M} by mroi v{__version__}"),
    ]
    for r, (label, value) in enumerate(details, start=top + 1):
        ws.cell(row=r, column=1, value=label).font = _font(bold=True)
        ws.cell(row=r, column=2, value=value).alignment = Alignment(wrap_text=True, vertical="top")
    _widths(ws, {"A": 22, "B": 110})
